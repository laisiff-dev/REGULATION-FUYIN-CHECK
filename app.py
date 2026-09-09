import os
import json
import glob
import openpyxl
from flask import Flask, render_template, jsonify, request, send_file
from audit_engine import run_compliance_audit, parse_docx

app = Flask(__name__)

DB_FILE = 'audit_results.json'
EXCEL_FILE = '115年法規清單檢核表_已完成檢核.xlsx'

def load_db():
    if not os.path.exists(DB_FILE):
        run_compliance_audit()
    with open(DB_FILE, 'r', encoding='utf-8') as f:
        return json.load(f)

def save_db(data):
    with open(DB_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/audit-results', methods=['GET'])
def get_audit_results():
    data = load_db()
    sheets = list(set(d['sheet'] for d in data))
    
    # Calculate statistics
    total = len(data)
    c3_pass = sum(1 for d in data if d['c3'] == 'V')
    c3_fail = sum(1 for d in data if d['c3'] == 'X')
    c4_proposals = sum(1 for d in data if d['c4'] == 'V')
    c5_proposals = sum(1 for d in data if d['c5'] == 'V')
    
    stats = {
        'total': total,
        'c3_pass': c3_pass,
        'c3_fail': c3_fail,
        'c3_pass_rate': round((c3_pass / total * 100), 1) if total else 0,
        'c4_proposals': c4_proposals,
        'c5_proposals': c5_proposals
    }
    
    return jsonify({
        'stats': stats,
        'sheets': sheets,
        'data': data
    })

@app.route('/api/update-item', methods=['POST'])
def update_item():
    req = request.json
    sheet_name = req.get('sheet')
    seq = req.get('seq')
    
    db = load_db()
    updated_item = None
    for item in db:
        if item['sheet'] == sheet_name and str(item['seq']) == str(seq):
            item['c1'] = req.get('c1', item['c1'])
            item['c2'] = req.get('c2', item['c2'])
            item['c3'] = req.get('c3', item['c3'])
            item['c4'] = req.get('c4', item['c4'])
            item['c5'] = req.get('c5', item['c5'])
            item['unit'] = req.get('unit', item['unit'])
            item['last_date'] = req.get('last_date', item['last_date'])
            item['notes'] = req.get('notes', item['notes'])
            updated_item = item
            break
            
    if updated_item:
        save_db(db)
        
        # Update Excel file as well
        if os.path.exists(EXCEL_FILE):
            wb = openpyxl.load_workbook(EXCEL_FILE)
            if sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                for r in range(4, sheet.max_row + 1):
                    if str(sheet.cell(r, 1).value) == str(seq):
                        sheet.cell(r, 4, updated_item['c1'])
                        sheet.cell(r, 5, updated_item['c2'])
                        sheet.cell(r, 6, updated_item['c3'])
                        sheet.cell(r, 7, updated_item['c4'])
                        sheet.cell(r, 8, updated_item['c5'])
                        sheet.cell(r, 9, updated_item['unit'])
                        sheet.cell(r, 10, updated_item['last_date'])
                        sheet.cell(r, 11, updated_item['notes'])
                        break
                wb.save(EXCEL_FILE)
                wb.save('115年法規清單檢核表.xlsx')
                
        return jsonify({'status': 'success', 'item': updated_item})
    return jsonify({'status': 'error', 'message': 'Item not found'}), 404

@app.route('/api/re-audit', methods=['POST'])
def re_audit():
    run_compliance_audit()
    return jsonify({'status': 'success', 'message': '重新檢核完成！'})

@app.route('/api/doc-detail', methods=['GET'])
def get_doc_detail():
    file_name = request.args.get('filename')
    if not file_name:
        return jsonify({'status': 'error', 'message': '未提供檔名'}), 400
        
    if not os.path.exists(file_name):
        # try adding 'x'
        if os.path.exists(file_name + 'x'):
            file_name = file_name + 'x'
        else:
            return jsonify({'status': 'error', 'message': f'找不到文件 {file_name}'}), 404
            
    try:
        doc_info = parse_docx(file_name)
        return jsonify({
            'status': 'success',
            'filename': file_name,
            'info': doc_info
        })
    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)}), 500

@app.route('/api/download-excel', methods=['GET'])
def download_excel():
    if os.path.exists(EXCEL_FILE):
        return send_file(
            EXCEL_FILE,
            as_attachment=True,
            download_name='115年法規清單檢核表_已完成檢核.xlsx'
        )
    return jsonify({'status': 'error', 'message': '報表檔案不存在'}), 404

if __name__ == '__main__':
    print("啟動 輔英科技大學 研發處法規檢核 Web 系統 [http://127.0.0.1:5000]")
    app.run(host='0.0.0.0', port=5000, debug=True)
